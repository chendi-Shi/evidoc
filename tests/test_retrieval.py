from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
from pathlib import Path

import pytest

from app.ingest import MAX_FILE_BYTES, extract_document
from app.retrieval import Store, tokenize


@pytest.fixture
def store(tmp_path):
    return Store(tmp_path / "knowledge.db")


def add(store, name="报销.md", text="上海住宿标准每晚 600 元。报销需要电子发票。", hash_value="first"):
    return store.add_document(name, [{"page": 1, "text": text}], hash_value)


def test_persistence_deduplication_and_cascade_delete(tmp_path):
    path = tmp_path / "knowledge.db"
    store = Store(path)
    doc = add(store)
    chunk = store.all_chunks()[0]
    store.store_embeddings({chunk["id"]: [1.0, 0.0]}, "test-model")
    duplicate = add(store, name="副本.md")
    assert duplicate["id"] == doc["id"] and duplicate["deduplicated"]
    reopened = Store(path)
    assert reopened.stats() == {"documents": 1, "chunks": 1}
    assert reopened.chunks_missing_embeddings("test-model") == []
    assert reopened.delete_document(doc["id"])
    assert not reopened.delete_document(doc["id"])
    assert reopened.list_documents() == []
    assert reopened.all_chunks() == []
    with reopened._connect() as connection:
        assert connection.execute("SELECT COUNT(*) FROM embeddings").fetchone()[0] == 0


def test_concurrent_identical_uploads_are_atomic(store):
    with ThreadPoolExecutor(max_workers=4) as pool:
        documents = list(pool.map(lambda _: add(store), range(8)))
    assert len({doc["id"] for doc in documents}) == 1
    assert store.stats()["documents"] == 1
    assert sum(not doc["deduplicated"] for doc in documents) == 1


def test_chinese_english_and_unknown_query(store):
    add(store)
    add(store, "发布.md", "Production deployment uses Docker and rollback steps.", "second")
    assert store.search("上海住宿上限", mode="bm25")["hits"][0]["name"] == "报销.md"
    assert store.search("Docker rollback", mode="bm25")["hits"][0]["name"] == "发布.md"
    assert store.search("quasar neutrino", mode="bm25")["hits"] == []
    assert "zh:报销" in tokenize("报销报销")


def test_parent_child_preserves_pdf_page_and_bounds(store):
    store.add_document("长文.pdf", [
        {"page": 2, "text": "第一页没有文字，本页是原 PDF 第二页。" * 120},
        {"page": 5, "text": "第五页记载独立的差旅标准。"},
    ], "long")
    chunks = store.all_chunks()
    assert {chunk["page"] for chunk in chunks} == {2, 5}
    assert all(len(chunk["text"]) <= 350 and len(chunk["parent_text"]) <= 1000 for chunk in chunks)
    assert all(chunk["text"] in chunk["parent_text"] for chunk in chunks)
    assert any(len(chunk["parent_text"]) > len(chunk["text"]) for chunk in chunks)


def test_document_filter_is_respected_by_both_retrievers(store):
    one = add(store)
    two = add(store, "其他.md", "住宿上限 450 元。", "second")
    chunks = store.all_chunks()
    store.store_embeddings({chunk["id"]: [1.0, 0.0] for chunk in chunks}, "test-model")
    for mode in ("bm25", "hybrid", "dense"):
        result = store.search("住宿", mode=mode, query_vector=[1, 0], embedding_model="test-model", document_ids=[two["id"]])
        assert result["hits"]
        assert {hit["document_id"] for hit in result["hits"]} == {two["id"]}
    assert store.search("住宿", document_ids=[])["hits"] == []
    assert one["id"] != two["id"]


def test_dense_fallback_is_explicit_and_wrong_dimensions_do_not_crash(store):
    add(store)
    chunk = store.all_chunks()[0]
    store.store_embeddings({chunk["id"]: [1, 0]}, "test-model")
    for kwargs in ({}, {"query_vector": [1, 0], "embedding_model": "unknown"},
                   {"query_vector": [1, 0, 0], "embedding_model": "test-model"},
                   {"query_vector": [0, 0], "embedding_model": "test-model"}):
        result = store.search("报销", mode="dense", **kwargs)
        assert result["mode_used"] == "bm25"
        assert result["hits"] and result["warnings"]


def test_partial_hybrid_keeps_sparse_documents(store):
    add(store, "语义.md", "员工休息制度。", "semantic")
    target = add(store, "报销.md", "报销材料必须包含电子发票。", "sparse")
    chunk = store.all_chunks()[0]
    store.store_embeddings({chunk["id"]: [1, 0]}, "test-model")
    result = store.search("报销电子发票", mode="hybrid", query_vector=[1, 0], embedding_model="test-model")
    assert result["mode_used"] == "hybrid"
    assert target["id"] in {hit["document_id"] for hit in result["hits"]}
    assert any("部分" in warning for warning in result["warnings"])
    assert len(store.chunks_missing_embeddings("test-model")) == 1


def test_dense_threshold_and_invalid_embeddings(store):
    add(store)
    chunk = store.all_chunks()[0]
    store.store_embeddings({chunk["id"]: [1, 0]}, "test-model")
    result = store.search("语义查询", mode="dense", query_vector=[0, 1], embedding_model="test-model")
    assert result["hits"] == [] and result["mode_used"] == "dense"
    assert any("未经校准" in warning for warning in result["warnings"])
    for vector in ([float("nan"), 0], [0, 0], [1, 0, 0]):
        with pytest.raises(ValueError):
            store.store_embeddings({chunk["id"]: vector}, "test-model")
    assert store.chunks_missing_embeddings("test-model") == []


def test_text_extraction_encodings_and_validation():
    text = "费用报销需要电子发票。"
    for encoding in ("utf-8", "utf-8-sig", "gb18030"):
        assert extract_document("说明.TXT", text.encode(encoding)) == [{"text": text, "page": 1}]
    for name, content in (("file.exe", b"abc"), ("file.md", b""), ("file.txt", b"  \n\t"),
                          ("file.txt", b"a\x00b"), ("file.pdf", b"bad pdf"),
                          ("file.txt", b"a" * (MAX_FILE_BYTES + 1))):
        with pytest.raises(ValueError):
            extract_document(name, content)


def test_docx_paragraphs_and_tables():
    from docx import Document

    document = Document()
    document.add_paragraph("年假 12 天")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "住宿标准"
    table.cell(0, 1).text = "600 元"
    buffer = BytesIO()
    document.save(buffer)
    pages = extract_document("制度.docx", buffer.getvalue())
    assert len(pages) == 1 and pages[0]["page"] == 1
    assert "年假 12 天" in pages[0]["text"] and "600 元" in pages[0]["text"]


def test_blank_and_encrypted_pdf_report_clear_errors():
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    buffer = BytesIO()
    writer.write(buffer)
    with pytest.raises(ValueError, match="OCR"):
        extract_document("scan.pdf", buffer.getvalue())
    writer.encrypt("secret")
    buffer = BytesIO()
    writer.write(buffer)
    with pytest.raises(ValueError, match="加密"):
        extract_document("encrypted.pdf", buffer.getvalue())


def test_synthetic_demo_answerable_documents_recalled(store):
    import hashlib
    import json

    root = Path(__file__).resolve().parents[1]
    for path in (root / "data" / "demo").glob("*.md"):
        content = path.read_bytes()
        store.add_document(path.name, extract_document(path.name, content), hashlib.sha256(content).hexdigest())
    questions = json.loads((root / "evaluation" / "questions.json").read_text(encoding="utf-8"))
    assert len(questions) == 12 and sum(item["unanswerable"] for item in questions) == 3
    for question in questions:
        if question["unanswerable"]:
            continue
        names = {hit["name"] for hit in store.search(question["question"], k=5, mode="bm25")["hits"]}
        assert set(question["relevant_documents"]) <= names, question["id"]
