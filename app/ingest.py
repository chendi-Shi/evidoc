"""Bounded, local text extraction. Image-only PDF pages require external OCR."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
import re
import zipfile


MAX_FILE_BYTES = 25 * 1024 * 1024
MAX_TEXT_CHARS = 2_000_000
MAX_PAGES = 500
SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".md", ".txt"}


def _clean(text: str) -> str:
    text = text.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[^\S\n]+", " ", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def extract_document(filename: str, content: bytes) -> list[dict]:
    """Return nonempty text pages, retaining original PDF page numbers.

    PDF extraction does not perform OCR. A mixed PDF retains searchable pages;
    image-only pages have no extractable text and are omitted.
    """
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise ValueError("不支持此文件格式；请上传 PDF、DOCX、Markdown 或 TXT。")
    if not content:
        raise ValueError("文件为空。")
    if len(content) > MAX_FILE_BYTES:
        raise ValueError("文件不能超过 25 MB。")
    pages: list[dict] = []
    if suffix == ".pdf":
        from pypdf import PdfReader

        try:
            reader = PdfReader(BytesIO(content))
            if reader.is_encrypted:
                raise ValueError("无法导入加密 PDF；请先提供未加密的副本。")
            if len(reader.pages) > MAX_PAGES:
                raise ValueError(f"PDF 不能超过 {MAX_PAGES} 页。")
            total = 0
            for index, page in enumerate(reader.pages, 1):
                text = _clean(page.extract_text() or "")
                total += len(text)
                if total > MAX_TEXT_CHARS:
                    raise ValueError("提取文本不能超过 200 万字符。")
                if text:
                    pages.append({"text": text, "page": index})
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError("无法解析 PDF；文件可能损坏或包含不支持的内容。") from exc
        if not pages:
            raise ValueError("PDF 没有可提取的文字。本项目不执行 OCR，请使用带文字层的 PDF。")
    elif suffix == ".docx":
        from docx import Document

        try:
            # Check the archive before python-docx expands its XML parts.
            with zipfile.ZipFile(BytesIO(content)) as archive:
                entries = archive.infolist()
                if len(entries) > 5000 or sum(item.file_size for item in entries) > 50 * 1024 * 1024:
                    raise ValueError("DOCX 解压后的内容过大。")
                if any(item.flag_bits & 1 for item in entries):
                    raise ValueError("无法导入加密 DOCX。")
            document = Document(BytesIO(content))
            paragraphs = [paragraph.text for paragraph in document.paragraphs]
            # Tables often contain the most important procedural facts.
            paragraphs.extend(
                " | ".join(cell.text for cell in row.cells)
                for table in document.tables for row in table.rows
            )
            pages = [{"text": _clean("\n".join(paragraphs)), "page": 1}]
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError("无法解析 DOCX；请确认文件为有效的 Word 文档。") from exc
    else:
        text = None
        for encoding in ("utf-8-sig", "gb18030"):
            try:
                text = content.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        if text is None or "\x00" in text:
            raise ValueError("文本文件编码不受支持；请另存为 UTF-8。")
        pages = [{"text": _clean(text), "page": 1}]
    pages = [page for page in pages if page["text"]]
    if not pages:
        raise ValueError("文档没有可检索的文字。")
    if sum(len(page["text"]) for page in pages) > MAX_TEXT_CHARS:
        raise ValueError("提取文本不能超过 200 万字符。")
    return pages
